"""文档解析器：把上传文件解析为带页码/章节信息的文本块（蓝图 §2.1 rag/parsers.py）。

支持范围（PRD-03 §3）：P0 集 pdf / docx / md / txt；P1 集 csv / xlsx；
image / other 仍抛 PARSE_UNSUPPORTED（PRD-06 §5.3）。

设计要点（为什么）：
- 输出统一为 `ParsedDoc.blocks`（text / page / section_title），后续切片器只
  面向这一种结构，不感知文件格式——新增格式时只需加一个解析分支。
- pdf 逐页解析并保留 1 基页码：PRD-06 §4.5 要求引用展示页码。
- docx 用段落样式识别标题（Heading/标题），md/txt 按 `#` 开头行切分：
  标题块开启新 block，正文累积进当前 block，使"章节"成为切片的天然边界。
- csv/xlsx 用首行表头作列名上下文，数据行展平为 "列名 值；列名 值" 句子：
  表格的行列结构对向量语义不友好，展平后每行成为独立可召回的事实命题；
  xlsx 每个 sheet 一个 block 并以 sheet 名作章节标题。
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field

from .errors import PARSE_EMPTY_TEXT, PARSE_UNSUPPORTED, RAGError

# markdown 标题行（# ~ ######），txt 里出现 `#` 行也按标题处理（无副作用）
_MD_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*\S)\s*$")

SUPPORTED_FILE_TYPES = ("pdf", "docx", "md", "txt", "csv", "xlsx")


@dataclass
class Block:
    """解析产出的最小结构块：一段正文 + 它的页码与所属章节标题。"""

    text: str
    page: int | None = None
    section_title: str | None = None


@dataclass
class ParsedDoc:
    """解析结果：有序块列表 + 解析出的纯文本总量（供判空与调试）。"""

    blocks: list[Block] = field(default_factory=list)

    @property
    def total_text(self) -> str:
        return "".join(b.text for b in self.blocks)


def parse_document(file_bytes: bytes, file_type: str) -> ParsedDoc:
    """按文件类型解析字节流为 ParsedDoc。

    失败语义（PRD-06 §5.3）：
    - 非支持格式 → RAGError(PARSE_UNSUPPORTED)
    - 解析后全文为空 → RAGError(PARSE_EMPTY_TEXT)（典型场景：扫描件 PDF）
    """
    kind = (file_type or "").lower().strip()
    if kind == "pdf":
        doc = _parse_pdf(file_bytes)
    elif kind == "docx":
        doc = _parse_docx(file_bytes)
    elif kind in ("md", "txt"):
        doc = _parse_markdown(file_bytes)
    elif kind == "csv":
        doc = _parse_csv(file_bytes)
    elif kind == "xlsx":
        doc = _parse_xlsx(file_bytes)
    else:
        raise RAGError(PARSE_UNSUPPORTED)
    if not doc.total_text.strip():
        raise RAGError(PARSE_EMPTY_TEXT)
    return doc


def _parse_pdf(file_bytes: bytes) -> ParsedDoc:
    """pypdf 逐页提取文本；每页一个 block，页码从 1 起。"""
    from pypdf import PdfReader  # 延迟导入：非 PDF 路径不付出导入成本

    try:
        reader = PdfReader(io.BytesIO(file_bytes))
    except Exception as exc:  # 损坏/加密 PDF 统一按"解析不出文本"处理
        raise RAGError(PARSE_EMPTY_TEXT, f"PDF 文件无法读取：{exc}") from exc
    blocks: list[Block] = []
    for index, page in enumerate(reader.pages):
        try:
            text = (page.extract_text() or "").strip()
        except Exception:
            text = ""  # 单页解析失败不拖垮整篇，空页由末尾总量判空兜底
        if text:
            blocks.append(Block(text=text, page=index + 1, section_title=None))
    return ParsedDoc(blocks=blocks)


def _parse_docx(file_bytes: bytes) -> ParsedDoc:
    """python-docx 解析：Heading/标题样式开启新章节块，正文段落累积进当前块。

    docx 没有可靠页码（分页由渲染决定），page 一律为 None——引用展示时
    按 PRD-06 §4.5 退化为展示章节标题。
    """
    import docx  # python-docx，延迟导入理由同上

    try:
        document = docx.Document(io.BytesIO(file_bytes))
    except Exception as exc:
        raise RAGError(PARSE_EMPTY_TEXT, f"Word 文件无法读取：{exc}") from exc
    blocks: list[Block] = []
    current_title: str | None = None
    buffer: list[str] = []

    def flush() -> None:
        text = "\n".join(buffer).strip()
        if text:
            blocks.append(Block(text=text, page=None, section_title=current_title))
        buffer.clear()

    for para in document.paragraphs:
        style_name = (para.style.name or "") if para.style is not None else ""
        text = para.text.strip()
        if style_name.lower().startswith("heading") or style_name.startswith("标题"):
            flush()
            current_title = text or current_title  # 空标题行不覆盖已有章节名
            continue
        if text:
            buffer.append(text)
    flush()
    return ParsedDoc(blocks=blocks)


def _parse_markdown(file_bytes: bytes) -> ParsedDoc:
    """md/txt 解析：按 markdown 标题切分章节块，正文原样累积。"""
    text = file_bytes.decode("utf-8", errors="replace")
    blocks: list[Block] = []
    current_title: str | None = None
    buffer: list[str] = []

    def flush() -> None:
        body = "\n".join(buffer).strip()
        if body:
            blocks.append(Block(text=body, page=None, section_title=current_title))
        buffer.clear()

    for line in text.splitlines():
        match = _MD_HEADING_RE.match(line)
        if match:
            flush()
            current_title = match.group(2).strip()
            continue
        buffer.append(line)
    flush()
    return ParsedDoc(blocks=blocks)


def _row_to_sentence(header: list[str], row: list) -> str:
    """把一行表格数据按表头列名展平为 "列名 值；列名 值" 句子（空值跳过）。"""
    pairs: list[str] = []
    for index, raw in enumerate(row):
        value = "" if raw is None else str(raw).strip()
        if not value:
            continue
        name = header[index] if index < len(header) and header[index] else f"列{index + 1}"
        pairs.append(f"{name} {value}")
    return "；".join(pairs) + "。" if pairs else ""


def _parse_csv(file_bytes: bytes) -> ParsedDoc:
    """csv 解析：首行表头作列名上下文，每个数据行一个 block（stdlib csv，零依赖）。

    编码用 utf-8-sig：Excel 导出的 csv 常带 BOM，不去掉会把 BOM 混进首个列名；
    errors="replace" 保证脏字节不阻断解析（乱码行由末尾总量判空兜底）。
    """
    text = file_bytes.decode("utf-8-sig", errors="replace")
    rows = [row for row in csv.reader(io.StringIO(text)) if any(cell.strip() for cell in row)]
    if not rows:
        return ParsedDoc(blocks=[])
    header = [cell.strip() for cell in rows[0]]
    blocks: list[Block] = []
    for row in rows[1:]:
        sentence = _row_to_sentence(header, row)
        if sentence:
            blocks.append(Block(text=sentence, page=None, section_title=None))
    return ParsedDoc(blocks=blocks)


def _parse_xlsx(file_bytes: bytes) -> ParsedDoc:
    """xlsx 解析（openpyxl）：每个 sheet 一个 block，sheet 名作章节标题；
    sheet 内首行表头作列名，数据行展平为句子逐行排列。

    无法读取时抛 PARSE_UNSUPPORTED 而非 PARSE_EMPTY_TEXT：xlsx 是 zip 容器，
    打不开通常意味着旧版 .xls 二进制格式或文件损坏，属于"格式不受支持"，
    引导用户转换格式比提示"扫描件"更准确（PRD-06 §5.3 的用户引导语义）。
    """
    try:
        from openpyxl import load_workbook  # 延迟导入：非 xlsx 路径不付出导入成本
    except ImportError as exc:  # pragma: no cover - openpyxl 在 pyproject 依赖中
        raise RAGError(PARSE_UNSUPPORTED, "缺少 openpyxl 依赖，无法解析 Excel 文件") from exc
    try:
        workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise RAGError(PARSE_UNSUPPORTED, "Excel 文件无法读取或格式不受支持，请确认为 .xlsx 文件") from exc
    try:
        blocks: list[Block] = []
        for sheet in workbook.worksheets:
            rows = [list(row) for row in sheet.iter_rows(values_only=True)]
            # 跳过整行全空的行，首行非空行作为表头
            rows = [row for row in rows if any(cell is not None and str(cell).strip() for cell in row)]
            if len(rows) < 2:
                continue  # 只有表头或全空的 sheet 没有可索引内容
            header = ["" if cell is None else str(cell).strip() for cell in rows[0]]
            sentences = [s for s in (_row_to_sentence(header, row) for row in rows[1:]) if s]
            if sentences:
                blocks.append(Block(text="\n".join(sentences), page=None, section_title=sheet.title))
        return ParsedDoc(blocks=blocks)
    finally:
        workbook.close()  # read_only 模式持有文件句柄，必须显式关闭
